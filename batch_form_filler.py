"""Offline batch fill for pre-existing *text* AcroForm PDF fields.

This additive tool does not redact, sanitize, flatten, or attest to the
correctness of a PDF. Each output must be reviewed in a PDF viewer.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path

from pypdf import PdfReader, PdfWriter


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def records(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if path.suffix.lower() == ".csv":
        content = None
        for encoding in ("utf-8-sig", "utf-8", "cp932"):
            try:
                content = path.read_text(encoding=encoding)
                break
            except UnicodeError:
                continue
        if content is None:
            raise ValueError("CSV encoding not supported")
        csv_rows = list(csv.reader(content.splitlines()))
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install openpyxl to use XLSX") from exc
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            csv_rows = list(wb.worksheets[0].iter_rows(values_only=True))
        finally:
            wb.close()
    else:
        raise ValueError("Only CSV / XLSX input is supported")
    if not csv_rows:
        raise ValueError("Empty input table")
    headers = [text(v) for v in csv_rows[0]]
    if not headers or not all(headers) or len(headers) != len(set(headers)):
        raise ValueError("Headers must be unique and nonempty")
    if len(csv_rows) > 5001:
        raise ValueError("MVP safety limit: 5000 records per run")
    result: list[dict[str, str]] = []
    for row_number, raw in enumerate(csv_rows[1:], start=2):
        if not any(text(cell) for cell in raw):
            continue
        if len(raw) > len(headers) and any(text(v) for v in raw[len(headers):]):
            raise ValueError(f"Unexpected extra columns in row {row_number}")
        vals = list(raw[:len(headers)]) + [""] * max(0, len(headers) - len(raw))
        result.append(dict(zip(headers, map(text, vals))))
    if not result:
        raise ValueError("No populated rows")
    return headers, result


def _form_fields(reader: PdfReader) -> list[str]:
    if reader.is_encrypted:
        raise ValueError("Encrypted PDF is not supported")
    root = reader.trailer["/Root"]
    # This form-filler is not a PDF sanitizer. Reject obvious active content
    # instead of passing it through into each generated document.
    if root.get("/OpenAction") or root.get("/AA"):
        raise ValueError("PDF contains document-level actions")
    names = root.get("/Names")
    if names:
        name_dict = names.get_object()
        if name_dict.get("/JavaScript") or name_dict.get("/EmbeddedFiles"):
            raise ValueError("PDF contains scripts or embedded files")
    for page in reader.pages:
        if page.get("/AA"):
            raise ValueError("PDF contains page-level actions")
        for ref in page.get("/Annots", []):
            annotation = ref.get_object()
            if annotation.get("/A") or annotation.get("/AA"):
                raise ValueError("PDF contains annotation actions")
    acro = root.get("/AcroForm")
    if acro and acro.get_object().get("/XFA"):
        raise ValueError("XFA forms are not supported")
    fields = reader.get_fields() or {}
    if not fields:
        raise ValueError("No AcroForm fields; this tool does not stamp non-fillable PDFs")
    for key, field in fields.items():
        typ = field.get("/FT")
        if typ == "/Sig":
            raise ValueError("Signed/signature field detected; aborting")
        if typ != "/Tx":
            raise ValueError(f"Unsupported field type for {key}: {typ} (text only)")
    return list(fields)


def batch_fill(template: str | Path, data: str | Path, output_parent: str | Path,
               mapping: dict[str, str] | None = None) -> Path:
    template = Path(template).expanduser().resolve()
    data = Path(data).expanduser().resolve()
    if not template.is_file() or not data.is_file():
        raise FileNotFoundError("Template/data file does not exist")
    if template == data:
        raise ValueError("Template and data must be distinct")
    header, rows = records(data)
    reader = PdfReader(template)
    names = _form_fields(reader)
    mapping = mapping or {field: field for field in names}
    if not mapping or any(name not in names for name in mapping):
        raise ValueError("Mapping contains unknown PDF field(s)")
    if any(col not in header for col in mapping.values()):
        raise ValueError("Mapping contains missing data column(s)")
    parent = Path(output_parent).expanduser().resolve()
    parent.mkdir(parents=True, exist_ok=True)
    run = Path(tempfile.mkdtemp(prefix="batch_pdf_forms_", dir=parent))
    summary: list[dict] = []
    try:
        for index, row in enumerate(rows, 1):
            values = {field: row[col] for field, col in mapping.items()}
            writer = PdfWriter()
            writer.append(reader)
            writer.update_page_form_field_values(None, values, auto_regenerate=False)
            out = run / f"filled_{index:04d}.pdf"
            # x mode avoids overwriting any output even if names collide
            with out.open("xb") as handle:
                writer.write(handle)
            verified_fields = PdfReader(out).get_fields() or {}
            mismatches = [key for key, val in values.items()
                          if text(verified_fields.get(key, {}).get("/V", "")) != val]
            unicode_review = any(any(ord(ch) > 127 for ch in value) for value in values.values())
            summary.append({
                "row_number": index + 1,
                "output": out.name,
                "sha256": sha256(out),
                "status": "REVIEW" if mismatches or unicode_review else "FIELD_VALUES_VERIFIED",
                "mismatched_fields": mismatches,
                "visual_review_required": True,
                "non_ascii_font_review": unicode_review,
            })
    except Exception:
        # A partial output is never marked successful; preserve it for inspection.
        (run / "INCOMPLETE.txt").write_text("Generation aborted; do not distribute PDFs in this folder.", encoding="utf-8")
        raise
    report = {
        "mode": "pdf_acroform_text_batch_mvp",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "template_sha256": sha256(template),
        "data_sha256": sha256(data),
        "mapping": mapping,
        "files": summary,
        "limitations": [
            "Text AcroForm only: no image-based PDF, XFA, signatures, checkboxes, or flattening.",
            "Field value verification is not visual verification.",
            "Non-ASCII text may require embedded fonts or manual font selection.",
            "Unmapped fields are left as-is; verify no sensitive template defaults remain.",
            "Not a security sanitizer; PDF documents may retain metadata or scripts.",
            "Review all outputs, especially non-ASCII Japanese text, before distribution.",
        ],
    }
    (run / "batch_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return run


def main() -> None:
    ap = argparse.ArgumentParser(description="Offline AcroForm batch fill; read-only source files")
    ap.add_argument("template", help="Fillable PDF template")
    ap.add_argument("data", help="CSV or XLSX with a header")
    ap.add_argument("output_parent")
    ap.add_argument("--mapping", help='JSON filename with {"pdf_field": "spreadsheet_column"}')
    args = ap.parse_args()
    mapping = None
    if args.mapping:
        mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8"))
        if not isinstance(mapping, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in mapping.items()):
            ap.error("Mapping must be a JSON object of strings")
    print(batch_fill(args.template, args.data, args.output_parent, mapping))


if __name__ == "__main__":
    main()
